# title="app/services/generator_service.py"
import os
import uuid
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from app.config import settings

class StructuredFileGeneratorService:
    @staticmethod
    def compile_pdf(markdown_content: str) -> str:
        """
        Translates raw structural text inputs into a printable PDF download binary.
        """
        filename = f"Generated_Report_{uuid.uuid4()}.pdf"
        dest_path = os.path.join(settings.OUTPUT_DIR, filename)
        
        doc = SimpleDocTemplate(dest_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        story = []
        styles = getSampleStyleSheet()
        
        # Build custom typography profiles
        custom_body_style = ParagraphStyle(
            'CustomBody',
            parent=styles['BodyText'],
            fontSize=11,
            leading=15,
            spaceAfter=10
        )
        
        lines = markdown_content.split('\n')
        for line in lines:
            cleaned = line.strip()
            if not cleaned:
                continue
            if cleaned.startswith('#'):
                # Handle structured markdown header tokens cleanly
                level = len(cleaned) - len(cleaned.lstrip('#'))
                text = cleaned.lstrip('#').strip()
                size = max(24 - (level * 3), 12)
                h_style = ParagraphStyle(f'H_{level}', parent=styles['Heading1'], fontSize=size, leading=size+4, spaceAfter=8, spaceBefore=12)
                story.append(Paragraph(text, h_style))
            else:
                story.append(Paragraph(cleaned, custom_body_style))
                
        doc.build(story)
        return filename

    @staticmethod
    def compile_docx(markdown_content: str) -> str:
        """
        Translates raw text payloads into formatted Microsoft Word OpenXML (.docx) asset profiles.
        """
        filename = f"Generated_Document_{uuid.uuid4()}.docx"
        dest_path = os.path.join(settings.OUTPUT_DIR, filename)
        
        doc = Document()
        lines = markdown_content.split('\n')
        
        for line in lines:
            cleaned = line.strip()
            if not cleaned:
                continue
            if cleaned.startswith('#'):
                level = len(cleaned) - len(cleaned.lstrip('#'))
                text = cleaned.lstrip('#').strip()
                doc.add_heading(text, level=min(level, 4))
            else:
                doc.add_paragraph(cleaned)
                
        doc.save(dest_path)
        return filename